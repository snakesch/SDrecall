#!/usr/bin/env python3
"""Add requested extraction evidence to the canonical FN diagnosis workbook."""

from __future__ import annotations

import argparse
import copy
import os
import re
import sys
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pysam


OBSERVATION_COLUMNS = [
    "Assembly",
    "FN allele",
    "Query SD region",
    "Homologous counterpart region",
    "RG index",
    "FC query pair index",
    "NFC/counterpart pair index",
    "Selected input BAM ALT-haplotype qnames",
    "Original raw BAM ALT-haplotype qnames",
    "ALT-haplotype depth / total depth at true ALT site",
]
OBSERVATION_AUDIT_COLUMN = "Realigned exact-allele AD distribution and FN cause"

QNAME_COLUMNS = [
    "Assembly",
    "FN allele",
    "ALT-haplotype qname",
    "Evidence source",
    "Selected input BAM alignment records",
    "Original raw BAM alignment records",
    "Realigned BAM alignment records",
]

NEW_COLUMNS = [
    "Should be extracted (FC/NFC pair)",
    "Present in extracted RG FASTQ",
]
QNAME_AUDIT_COLUMN = "Realigned read pair produces exact FN ALT"

PIPELINE_ROOTS = {
    "hg38": Path(
        "/paedyl01/disk1/yangyxt/SDrecall-test/results/"
        "sdrecall_rust_avg50x_20260707/hg38/"
        "HG002_hg38_exome_avg50x_SDrecall"
    ),
    "hg19": Path(
        "/paedyl01/disk1/yangyxt/SDrecall-test/results/"
        "sdrecall_rust_avg50x_20260707_hg19_ucscfix/hg19/"
        "HG002_hg19_exome_avg50x_SDrecall"
    ),
    "chm13": Path(
        "/paedyl01/disk1/yangyxt/SDrecall-test/results/"
        "sdrecall_rust_avg50x_20260707/t2t/"
        "HG002_chm13_exome_avg50x_SDrecall"
    ),
}

INPUT_SAMS = {
    assembly: Path(
        "/paedyl01/disk1/yangyxt/SDrecall-rust-migration/"
        f"test_tmp/fn_qname_alignment_audit_20260721/{assembly}.all_input_records.sam"
    )
    for assembly in PIPELINE_ROOTS
}

PAIR_RE = re.compile(
    r"(?P<rg>RG\d+)\s*\|\s*(?P<kind>FC|NFC):(?P<pair>RG\d+_\d+)\s*\|\s*"
    r"(?P<chrom>[^:;\s]+):(?P<start>\d+)-(?P<end>\d+)"
)
PAIR_TAG_RE = re.compile(r"^(?P<kind>FC|NFC):(?P<pair>RG\d+_\d+)$")
RG_RE = re.compile(r"^RG\d+$")
CIGAR_RE = re.compile(r"(\d+)([MIDNSHP=X])")
QNAME_RG_SUFFIX_RE = re.compile(r":RG\d+$")
QNAME_MATE_SUFFIX_RE = re.compile(r"/[12]$")
REF_CONSUMING = frozenset("MDN=X")


@dataclass(frozen=True)
class Interval:
    chrom: str
    start: int
    end: int

    def overlaps(self, other: "Interval") -> bool:
        return self.chrom == other.chrom and self.start < other.end and other.start < self.end


@dataclass(frozen=True)
class PairRegion:
    kind: str
    rg: str
    pair: str
    displayed_interval: Interval

    @property
    def label(self) -> str:
        return f"{self.kind}:{self.pair}"


@dataclass
class SamRecord:
    qname: str
    flag: int
    interval: Interval | None
    mate_chrom: str | None
    mate_start: int | None
    sequence_length: int
    tags: dict[str, str]

    @property
    def mate_number(self) -> int:
        if self.flag & 0x40:
            return 1
        if self.flag & 0x80:
            return 2
        return 0

    @property
    def is_unmapped(self) -> bool:
        return bool(self.flag & 0x4)

    @property
    def is_mate_unmapped(self) -> bool:
        return bool(self.flag & 0x8)

    @property
    def is_paired(self) -> bool:
        return bool(self.flag & 0x1)

    def passes_nfc_filter(self) -> bool:
        if "SA" in self.tags:
            return False
        if "XA" in self.tags:
            return True
        try:
            return abs(int(self.tags["AS"]) - int(self.tags["XS"])) < 10
        except (KeyError, ValueError):
            return False


def worksheet_headers(ws: openpyxl.worksheet.worksheet.Worksheet) -> list[str]:
    return [cell.value for cell in ws[1]]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--workbook",
        type=Path,
        default=Path(
            "/paedyl01/disk1/yangyxt/SDrecall-rust-migration/"
            "FN_CAUSE_DIAGNOSIS.xlsx"
        ),
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="calculate and report evidence without changing the workbook",
    )
    parser.add_argument(
        "--pipeline-root",
        action="append",
        default=[],
        metavar="ASSEMBLY=PATH",
        help="override a production run root; may be supplied once per assembly",
    )
    parser.add_argument(
        "--input-sam",
        action="append",
        default=[],
        metavar="ASSEMBLY=PATH",
        help="override an input-alignment SAM audit file; may be supplied once per assembly",
    )
    return parser.parse_args()


def apply_path_overrides(
    paths: dict[str, Path], assignments: list[str], option: str
) -> None:
    for assignment in assignments:
        assembly, separator, value = assignment.partition("=")
        if not separator or not assembly or not value:
            raise ValueError(f"{option} expects ASSEMBLY=PATH, got {assignment!r}")
        if assembly not in paths:
            raise ValueError(
                f"{option} has unknown assembly {assembly!r}; "
                f"expected one of {', '.join(sorted(paths))}"
            )
        paths[assembly] = Path(value)


def canonical_qname(name: str) -> str:
    name = name.strip()
    if name.startswith(("@", ">")):
        name = name[1:]
    name = name.split(None, 1)[0]
    name = QNAME_MATE_SUFFIX_RE.sub("", name)
    return QNAME_RG_SUFFIX_RE.sub("", name)


def cigar_reference_length(cigar: str) -> int:
    if cigar == "*":
        return 0
    return sum(int(length) for length, op in CIGAR_RE.findall(cigar) if op in REF_CONSUMING)


def parse_sam(path: Path) -> dict[str, list[SamRecord]]:
    records: dict[str, list[SamRecord]] = defaultdict(list)
    with path.open() as handle:
        for line in handle:
            if not line or line.startswith("@"):
                continue
            fields = line.rstrip("\n").split("\t")
            if len(fields) < 11:
                raise ValueError(f"malformed SAM record in {path}: {line[:100]!r}")
            qname = canonical_qname(fields[0])
            flag = int(fields[1])
            chrom = fields[2]
            pos = int(fields[3])
            cigar = fields[5]
            mate_chrom = fields[6]
            mate_pos = int(fields[7])
            interval = None
            if chrom != "*" and pos > 0 and not (flag & 0x4):
                start = pos - 1
                interval = Interval(chrom, start, start + cigar_reference_length(cigar))
            tags: dict[str, str] = {}
            for field in fields[11:]:
                parts = field.split(":", 2)
                if len(parts) == 3:
                    tags[parts[0]] = parts[2]
            if mate_chrom == "=":
                mate_chrom = chrom
            if mate_chrom == "*" or mate_pos <= 0:
                mate_chrom = None
                mate_start = None
            else:
                mate_start = mate_pos - 1
            records[qname].append(
                SamRecord(
                    qname,
                    flag,
                    interval,
                    mate_chrom,
                    mate_start,
                    0 if fields[9] == "*" else len(fields[9]),
                    tags,
                )
            )
    return records


def parse_bed(path: Path) -> list[Interval]:
    intervals: list[Interval] = []
    with path.open() as handle:
        for line in handle:
            if not line.strip() or line.startswith("#"):
                continue
            fields = line.rstrip("\n").split("\t")
            intervals.append(Interval(fields[0], int(fields[1]), int(fields[2])))
    return intervals


def parse_pair_regions(value: str, expected_kind: str) -> list[PairRegion]:
    regions: list[PairRegion] = []
    seen: set[tuple[str, str, str, int, int]] = set()
    for match in PAIR_RE.finditer(value or ""):
        if match.group("kind") != expected_kind:
            continue
        key = (
            match.group("rg"),
            match.group("pair"),
            match.group("chrom"),
            int(match.group("start")),
            int(match.group("end")),
        )
        if key in seen:
            continue
        seen.add(key)
        regions.append(
            PairRegion(
                expected_kind,
                match.group("rg"),
                match.group("pair"),
                Interval(
                    match.group("chrom"),
                    int(match.group("start")) - 1,
                    int(match.group("end")),
                ),
            )
        )
    return regions


def parse_tagged_pair_regions(path: Path) -> list[PairRegion]:
    regions: list[PairRegion] = []
    with path.open() as handle:
        for line in handle:
            if not line.strip() or line.startswith("#"):
                continue
            fields = line.rstrip("\n").split("\t")
            if len(fields) < 7:
                continue
            match = PAIR_TAG_RE.match(fields[6])
            if match is None:
                continue
            pair = match.group("pair")
            regions.append(
                PairRegion(
                    match.group("kind"),
                    pair.split("_", 1)[0],
                    pair,
                    Interval(fields[0], int(fields[1]), int(fields[2])),
                )
            )
    return regions


def production_extraction_intervals(
    assembly: str,
    rg: str,
    kind: str,
    bed_cache: dict[Path, list[Interval]],
) -> list[Interval]:
    rg_dir = PIPELINE_ROOTS[assembly] / "realign_groups" / rg
    if kind == "FC":
        bed_path = rg_dir / f"{rg}.fc_target.bed"
    else:
        bed_path = rg_dir / f"{rg}.counterparts_regions.targeted.bed"
    if not bed_path.is_file():
        raise FileNotFoundError(f"missing production {kind} extraction BED: {bed_path}")
    return bed_cache.setdefault(bed_path, parse_bed(bed_path))


def rg_source_regions(
    assembly: str,
    rg: str,
    source_cache: dict[Path, list[PairRegion]],
) -> list[PairRegion]:
    path = (
        PIPELINE_ROOTS[assembly]
        / "realign_groups"
        / rg
        / f"{rg}_related_homo_regions.bed"
    )
    if not path.is_file():
        raise FileNotFoundError(f"missing RG source-region BED: {path}")
    return source_cache.setdefault(path, parse_tagged_pair_regions(path))


def merged_intervals(intervals: list[Interval]) -> list[Interval]:
    merged: list[Interval] = []
    for interval in sorted(intervals, key=lambda item: (item.chrom, item.start, item.end)):
        if merged and merged[-1].chrom == interval.chrom and interval.start <= merged[-1].end:
            previous = merged[-1]
            merged[-1] = Interval(previous.chrom, previous.start, max(previous.end, interval.end))
        else:
            merged.append(interval)
    return merged


def extracted_pair_for_interval(
    records: list[SamRecord], extraction_interval: Interval
) -> tuple[SamRecord, SamRecord] | None:
    mates: dict[int, SamRecord] = {}
    for record in records:
        if record.interval is not None and record.interval.overlaps(extraction_interval):
            if record.mate_number in (1, 2):
                mates[record.mate_number] = record
    if 1 in mates and 2 in mates:
        return mates[1], mates[2]
    if len(mates) != 1:
        return None

    anchor = next(iter(mates.values()))
    if (
        not anchor.is_paired
        or anchor.is_mate_unmapped
        or anchor.mate_chrom is None
        or anchor.mate_start is None
    ):
        return None
    mate_window = Interval(
        anchor.mate_chrom,
        max(anchor.mate_start - 5, 0),
        anchor.mate_start + 5 + anchor.sequence_length,
    )
    for candidate in records:
        if (
            candidate.mate_number != anchor.mate_number
            and candidate.interval is not None
            and candidate.interval.overlaps(mate_window)
        ):
            mates[candidate.mate_number] = candidate
            break
    if 1 in mates and 2 in mates:
        return mates[1], mates[2]
    return None


def extraction_labels(
    pair: tuple[SamRecord, SamRecord],
    extraction_interval: Interval,
    kind: str,
    source_regions: list[PairRegion],
) -> set[str]:
    overlapping_records = [
        record.interval
        for record in pair
        if record.interval is not None and record.interval.overlaps(extraction_interval)
    ]
    labels = {
        region.label
        for region in source_regions
        if region.kind == kind
        and any(
            record_interval.overlaps(region.displayed_interval)
            for record_interval in overlapping_records
        )
    }
    if labels:
        return labels
    # A read can be fetched from the 500-bp FC padding without touching the
    # unpadded FC itself. In that case, identify the subgroup(s) that generated
    # the production extraction interval rather than dropping the pathway label.
    return {
        region.label
        for region in source_regions
        if region.kind == kind
        and extraction_interval.overlaps(region.displayed_interval)
    }


def extraction_outcome(
    records: list[SamRecord],
    intervals: list[Interval],
    kind: str,
    source_regions: list[PairRegion],
) -> tuple[str, set[str]]:
    saw_overlap = False
    saw_complete_pair = False
    eligible = False
    labels: set[str] = set()
    for interval in merged_intervals(intervals):
        if not any(
            record.interval is not None and record.interval.overlaps(interval)
            for record in records
        ):
            continue
        saw_overlap = True
        pair = extracted_pair_for_interval(records, interval)
        if pair is None:
            continue
        saw_complete_pair = True
        if kind == "FC" or any(record.passes_nfc_filter() for record in pair):
            eligible = True
            labels.update(extraction_labels(pair, interval, kind, source_regions))
    if eligible:
        return "eligible", labels
    if not saw_overlap:
        return "not overlapped", set()
    if not saw_complete_pair:
        return "overlapped; incomplete pair", set()
    return "overlapped; NFC filter failed", set()


def extraction_status(
    assembly: str,
    qname: str,
    rgs: list[str],
    input_records: dict[str, list[SamRecord]],
    bed_cache: dict[Path, list[Interval]],
    source_cache: dict[Path, list[PairRegion]],
) -> str:
    records = input_records.get(qname, [])
    if not records:
        return "No | absent from selected input BAM"

    eligible_paths: list[str] = []
    ineligible_paths: list[str] = []
    for rg in rgs:
        source_regions = rg_source_regions(assembly, rg, source_cache)
        for kind in ("FC", "NFC"):
            intervals = production_extraction_intervals(
                assembly, rg, kind, bed_cache
            )
            outcome, labels = extraction_outcome(
                records, intervals, kind, source_regions
            )
            if outcome == "eligible":
                detail = ",".join(sorted(labels)) or f"{kind}:{rg} aggregate"
                eligible_paths.append(f"Yes | {detail}")
            else:
                ineligible_paths.append(f"{rg} {kind} {outcome}")

    if eligible_paths:
        return "; ".join(dict.fromkeys(eligible_paths))
    detail = "; ".join(ineligible_paths) if ineligible_paths else "no FC/NFC pair"
    return f"No | {detail}"


def scan_fastq_qnames(path: Path) -> set[str]:
    names: set[str] = set()
    with path.open() as handle:
        line_number = 0
        for line in handle:
            line_number += 1
            if line_number % 4 == 1:
                names.add(canonical_qname(line))
    if line_number % 4:
        raise ValueError(f"truncated FASTQ: {path}")
    return names


def load_fastq_membership(
    needed: dict[tuple[str, str], set[str]],
) -> dict[tuple[str, str, str], tuple[bool, bool]]:
    membership: dict[tuple[str, str, str], tuple[bool, bool]] = {}
    for (assembly, rg), qnames in sorted(needed.items()):
        recall_dir = PIPELINE_ROOTS[assembly] / "recall_results"
        prefix = recall_dir / f"HG002.sdrecall.only_{rg}"
        r1_path = Path(f"{prefix}.r1.fastq")
        r2_path = Path(f"{prefix}.r2.fastq")
        if not r1_path.is_file() or not r2_path.is_file():
            raise FileNotFoundError(f"missing extracted FASTQ pair: {r1_path}, {r2_path}")
        r1_names = scan_fastq_qnames(r1_path)
        r2_names = scan_fastq_qnames(r2_path)
        for qname in qnames:
            membership[(assembly, rg, qname)] = (qname in r1_names, qname in r2_names)
        print(
            f"FASTQ {assembly} {rg}: R1={len(r1_names)} R2={len(r2_names)} "
            f"audited_qnames={len(qnames)}",
            file=sys.stderr,
        )
    return membership


def scan_realigned_bams(
    needed: dict[tuple[str, str], set[str]],
) -> dict[tuple[str, str, str], list[str]]:
    records: dict[tuple[str, str, str], list[str]] = defaultdict(list)
    for (assembly, rg), qnames in sorted(needed.items()):
        bam_path = (
            PIPELINE_ROOTS[assembly]
            / "recall_results"
            / f"HG002.sdrecall.only_{rg}.raw.bam"
        )
        if not bam_path.is_file():
            raise FileNotFoundError(f"missing realigned BAM: {bam_path}")
        found: set[str] = set()
        with pysam.AlignmentFile(bam_path, "rb", check_sq=False) as bam:
            for record in bam.fetch(until_eof=True):
                qname = canonical_qname(record.query_name or "")
                if qname not in qnames:
                    continue
                found.add(qname)
                if record.is_unmapped or record.reference_name is None:
                    location = "*"
                    cigar = "*"
                else:
                    location = f"{record.reference_name}:{record.reference_start + 1}"
                    cigar = record.cigarstring or "*"
                text = f"{rg} | {location} | {cigar}"
                key = (assembly, rg, qname)
                if text not in records[key]:
                    records[key].append(text)
        print(
            f"BAM {assembly} {rg}: audited_qnames={len(qnames)} "
            f"found_qnames={len(found)}",
            file=sys.stderr,
        )
    return records


def fastq_status(
    assembly: str,
    qname: str,
    rgs: list[str],
    membership: dict[tuple[str, str, str], tuple[bool, bool]],
) -> str:
    present: list[str] = []
    for rg in rgs:
        in_r1, in_r2 = membership[(assembly, rg, qname)]
        if in_r1 and in_r2:
            present.append(f"Yes | {rg} | R1+R2")
        elif in_r1 or in_r2:
            mate = "R1" if in_r1 else "R2"
            present.append(f"Partial | {rg} | {mate} only")
    return "; ".join(present) if present else "No"


def realigned_status(
    assembly: str,
    qname: str,
    rgs: list[str],
    records: dict[tuple[str, str, str], list[str]],
) -> str:
    values: list[str] = []
    for rg in rgs:
        values.extend(records.get((assembly, rg, qname), []))
    return "\n".join(values)


def copy_cell_style(source: openpyxl.cell.cell.Cell, target: openpyxl.cell.cell.Cell) -> None:
    if source.has_style:
        target._style = copy.copy(source._style)
    target.font = copy.copy(source.font)
    target.fill = copy.copy(source.fill)
    target.border = copy.copy(source.border)
    target.alignment = copy.copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy.copy(source.protection)


def main() -> None:
    args = parse_args()
    apply_path_overrides(PIPELINE_ROOTS, args.pipeline_root, "--pipeline-root")
    apply_path_overrides(INPUT_SAMS, args.input_sam, "--input-sam")
    workbook = openpyxl.load_workbook(args.workbook)
    if workbook.sheetnames != ["FN observations", "Qname mappings"]:
        raise ValueError(f"unexpected workbook sheets: {workbook.sheetnames}")

    observations = workbook["FN observations"]
    mappings = workbook["Qname mappings"]
    observation_headers = worksheet_headers(observations)
    if observation_headers not in (
        OBSERVATION_COLUMNS,
        OBSERVATION_COLUMNS + [OBSERVATION_AUDIT_COLUMN],
    ):
        raise ValueError("FN observations schema differs from the agreed canonical format")
    mapping_headers = worksheet_headers(mappings)
    if mapping_headers not in (
        QNAME_COLUMNS,
        QNAME_COLUMNS + NEW_COLUMNS,
        QNAME_COLUMNS + NEW_COLUMNS + [QNAME_AUDIT_COLUMN],
    ):
        raise ValueError("Qname mappings schema differs from the agreed canonical format")

    observation_rows = list(observations.iter_rows(min_row=2, values_only=True))
    observation_by_key = {
        (str(row[0]), str(row[1])): row for row in observation_rows
    }
    mapping_rows = list(mappings.iter_rows(min_row=2, max_col=7, values_only=True))

    input_records = {assembly: parse_sam(path) for assembly, path in INPUT_SAMS.items()}
    rgs_by_fn: dict[tuple[str, str], list[str]] = {}
    needed_fastq_qnames: dict[tuple[str, str], set[str]] = defaultdict(set)
    for key, observation in observation_by_key.items():
        regions = parse_pair_regions(str(observation[2]), "FC")
        regions.extend(parse_pair_regions(str(observation[3]), "NFC"))
        rgs = sorted({region.rg for region in regions}, key=lambda rg: int(rg[2:]))
        if not rgs:
            rgs = sorted(
                {rg.strip() for rg in str(observation[4]).split(";") if RG_RE.match(rg.strip())},
                key=lambda rg: int(rg[2:]),
            )
        rgs_by_fn[key] = rgs

    for row in mapping_rows:
        assembly, fn, qname = str(row[0]), str(row[1]), canonical_qname(str(row[2]))
        for rg in rgs_by_fn[(assembly, fn)]:
            needed_fastq_qnames[(assembly, rg)].add(qname)

    membership = load_fastq_membership(needed_fastq_qnames)
    realigned_records = scan_realigned_bams(needed_fastq_qnames)
    bed_cache: dict[Path, list[Interval]] = {}
    source_cache: dict[Path, list[PairRegion]] = {}
    calculated: list[tuple[str, str, str]] = []
    for row in mapping_rows:
        assembly, fn, qname = str(row[0]), str(row[1]), canonical_qname(str(row[2]))
        key = (assembly, fn)
        calculated.append(
            (
                extraction_status(
                    assembly,
                    qname,
                    rgs_by_fn[key],
                    input_records[assembly],
                    bed_cache,
                    source_cache,
                ),
                fastq_status(assembly, qname, rgs_by_fn[key], membership),
                realigned_status(
                    assembly, qname, rgs_by_fn[key], realigned_records
                ),
            )
        )

    print(f"workbook={args.workbook}")
    print(f"FN_observation_rows={len(observation_rows)}")
    print(f"qname_mapping_rows={len(mapping_rows)}")
    print(f"unique_extraction_BEDs={len(bed_cache)}")
    print(f"unique_RG_source_BEDs={len(source_cache)}")
    should_yes = ["Yes |" in result[0] for result in calculated]
    fastq_yes = ["Yes |" in result[1] for result in calculated]
    realigned_yes = [bool(result[2]) for result in calculated]
    print(f"should_yes_rows={sum(should_yes)}")
    print(f"FASTQ_yes_rows={sum(fastq_yes)}")
    print(
        "FASTQ_partial_rows="
        f"{sum(result[1].startswith('Partial') for result in calculated)}"
    )
    print(f"realigned_BAM_yes_rows={sum(realigned_yes)}")
    print(
        "FASTQ_yes_realigned_BAM_no_rows="
        f"{sum(observed and not realigned for observed, realigned in zip(fastq_yes, realigned_yes))}"
    )
    missing_after_eligibility = [
        (row, result)
        for row, result, expected, observed in zip(
            mapping_rows, calculated, should_yes, fastq_yes
        )
        if expected and not observed
    ]
    print(f"should_yes_FASTQ_no_rows={len(missing_after_eligibility)}")
    for row, result in missing_after_eligibility[:20]:
        print(
            "should_yes_FASTQ_no\t"
            f"{row[0]}\t{row[1]}\t{row[2]}\t{result[0]}\t{result[1]}"
        )

    target_key = ("hg19", "chr1:207743256:G:T")
    target = [
        (row, result)
        for row, result in zip(mapping_rows, calculated)
        if (str(row[0]), str(row[1])) == target_key
    ]
    print(f"target_rows={len(target)}")
    print(f"target_input_rows={sum(str(row[3]) == 'input+raw' for row, _ in target)}")
    print(f"target_should_yes={sum('Yes |' in result[0] for _, result in target)}")
    print(f"target_FASTQ_yes={sum('Yes |' in result[1] for _, result in target)}")

    if args.dry_run:
        return

    for row_number, result in enumerate(calculated, start=2):
        mappings.cell(row_number, 7, result[2])

    for offset, header in enumerate(NEW_COLUMNS, start=8):
        header_cell = mappings.cell(1, offset, header)
        copy_cell_style(mappings.cell(1, 7), header_cell)
        for row_number, values in enumerate(calculated, start=2):
            cell = mappings.cell(row_number, offset, values[offset - 8])
            copy_cell_style(mappings.cell(row_number, 7), cell)
            cell.alignment = copy.copy(mappings.cell(row_number, 7).alignment)

    mappings.column_dimensions["H"].width = 48
    mappings.column_dimensions["I"].width = 30
    if mappings.auto_filter.ref:
        last_column = "J" if mappings.max_column == 10 else "I"
        mappings.auto_filter.ref = f"A1:{last_column}{mappings.max_row}"

    if worksheet_headers(observations) != observation_headers:
        raise AssertionError("FN observations changed during workbook update")
    if list(observations.iter_rows(min_row=2, values_only=True)) != observation_rows:
        raise AssertionError("FN observation values changed during workbook update")

    tmp_path = args.workbook.with_name(f".{args.workbook.name}.{os.getpid()}.tmp")
    workbook.save(tmp_path)
    os.replace(tmp_path, args.workbook)
    print("saved=true")


if __name__ == "__main__":
    main()
