#!/usr/bin/env python3
"""Refresh run-dependent fields in the canonical FN diagnosis workbook.

The workbook format is intentionally fixed. This script updates only the
physical FC/NFC labels, RG/pair indices, and strict ALT-haplotype depth after
a rerun whose exact FN set is unchanged. Qname evidence is left untouched for
``update_fn_qname_extraction_audit.py`` to refresh against the new BAM/FASTQ
outputs.
"""

from __future__ import annotations

import argparse
import csv
import os
import re
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

import openpyxl


OBSERVATION_COLUMNS = [
    "Assembly",
    "FN allele",
    "Query SD region",
    "Homologous counterpart region",
    "RG index",
    "FC query pair index",
    "NFC/counterpart pair index",
    "Selected input BAM ALT-haplotype qnames",
    "Original raw BAM ALT-haplotype qnames",
    "ALT-haplotype depth / total depth at true ALT site",
]
OBSERVATION_AUDIT_COLUMN = "Realigned exact-allele AD distribution and FN cause"

QNAME_COLUMNS = [
    "Assembly",
    "FN allele",
    "ALT-haplotype qname",
    "Evidence source",
    "Selected input BAM alignment records",
    "Original raw BAM alignment records",
    "Realigned BAM alignment records",
    "Should be extracted (FC/NFC pair)",
    "Present in extracted RG FASTQ",
]
QNAME_AUDIT_COLUMN = "Realigned read pair produces exact FN ALT"

OLD_PIPELINE_ROOTS = {
    "hg38": Path(
        "/paedyl01/disk1/yangyxt/SDrecall-test/results/"
        "sdrecall_rust_avg50x_20260707/hg38/"
        "HG002_hg38_exome_avg50x_SDrecall"
    ),
    "hg19": Path(
        "/paedyl01/disk1/yangyxt/SDrecall-test/results/"
        "sdrecall_rust_avg50x_20260707_hg19_ucscfix/hg19/"
        "HG002_hg19_exome_avg50x_SDrecall"
    ),
    "chm13": Path(
        "/paedyl01/disk1/yangyxt/SDrecall-test/results/"
        "sdrecall_rust_avg50x_20260707/t2t/"
        "HG002_chm13_exome_avg50x_SDrecall"
    ),
}

TAG_RE = re.compile(r"^(?P<kind>FC|NFC):(?P<pair>RG\d+_\d+)$")
DISPLAY_RE = re.compile(
    r"^(?P<rg>RG\d+) \| (?P<kind>FC|NFC):(?P<pair>RG\d+_\d+) \| "
    r"(?P<chrom>[^:]+):(?P<start>\d+)-(?P<end>\d+)\((?P<strand>[+-])\)$"
)


@dataclass(frozen=True)
class Region:
    kind: str
    rg: str
    pair: str
    chrom: str
    start: int
    end: int
    strand: str

    @property
    def coord_key(self) -> tuple[str, int, int, str]:
        return self.chrom, self.start, self.end, self.strand

    @property
    def label(self) -> str:
        return f"{self.kind}:{self.pair}"

    def display(self) -> str:
        return (
            f"{self.rg} | {self.label} | "
            f"{self.chrom}:{self.start + 1}-{self.end}({self.strand})"
        )


@dataclass
class RegionCatalog:
    by_label: dict[str, Region]
    fc_by_coord: dict[tuple[str, int, int, str], list[Region]]
    by_pair_kind_coord: dict[
        tuple[str, str, tuple[str, int, int, str]], list[Region]
    ]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument(
        "--new-pipeline-root",
        action="append",
        required=True,
        metavar="ASSEMBLY=PATH",
    )
    parser.add_argument(
        "--haplotype-dir",
        action="append",
        required=True,
        metavar="ASSEMBLY=PATH",
        help="directory containing target_rg_*.tsv and pre_fp_raw_per_rg_summary.tsv",
    )
    parser.add_argument(
        "--old-pipeline-root",
        action="append",
        default=[],
        metavar="ASSEMBLY=PATH",
    )
    parser.add_argument("--dry-run", action="store_true")
    return parser.parse_args()


def parse_assignments(assignments: list[str], option: str) -> dict[str, Path]:
    parsed: dict[str, Path] = {}
    for assignment in assignments:
        assembly, separator, value = assignment.partition("=")
        if not separator or not assembly or not value:
            raise ValueError(f"{option} expects ASSEMBLY=PATH, got {assignment!r}")
        if assembly in parsed:
            raise ValueError(f"{option} repeats assembly {assembly!r}")
        parsed[assembly] = Path(value)
    return parsed


def rg_sort_key(rg: str) -> tuple[int, str]:
    match = re.fullmatch(r"RG(\d+)", rg)
    return (int(match.group(1)), rg) if match else (10**9, rg)


def read_catalog(root: Path) -> RegionCatalog:
    paths = sorted((root / "realign_groups").glob("RG*/RG*_related_homo_regions.bed"))
    if not paths:
        raise FileNotFoundError(f"no RG source-region BEDs under {root}")

    by_label: dict[str, Region] = {}
    fc_by_coord: dict[tuple[str, int, int, str], list[Region]] = defaultdict(list)
    by_pair_kind_coord: dict[
        tuple[str, str, tuple[str, int, int, str]], list[Region]
    ] = defaultdict(list)
    for path in paths:
        rg = path.parent.name
        with path.open() as handle:
            for line in handle:
                if not line.strip() or line.startswith("#"):
                    continue
                fields = line.rstrip("\n").split("\t")
                if len(fields) < 7:
                    continue
                match = TAG_RE.fullmatch(fields[6])
                if match is None:
                    continue
                pair = match.group("pair")
                region = Region(
                    kind=match.group("kind"),
                    rg=rg,
                    pair=pair,
                    chrom=fields[0],
                    start=int(fields[1]),
                    end=int(fields[2]),
                    strand=fields[5],
                )
                if region.kind == "FC":
                    previous = by_label.setdefault(region.label, region)
                    if previous != region:
                        raise ValueError(f"conflicting FC label {region.label} under {root}")
                    fc_by_coord[region.coord_key].append(region)
                by_pair_kind_coord[(pair, region.kind, region.coord_key)].append(region)
    return RegionCatalog(dict(by_label), dict(fc_by_coord), dict(by_pair_kind_coord))


def parse_displayed_region(line: str, catalog: RegionCatalog) -> Region:
    match = DISPLAY_RE.fullmatch(line)
    if match is None:
        raise ValueError(f"unexpected canonical region display: {line!r}")
    label = f"FC:{match.group('pair')}"
    owner = catalog.by_label.get(label)
    if owner is None:
        raise ValueError(f"old pair owner {label} is absent from its pipeline root")
    expected = (
        match.group("chrom"),
        int(match.group("start")) - 1,
        int(match.group("end")),
        match.group("strand"),
    )
    if match.group("kind") == "FC" and owner.coord_key != expected:
        raise ValueError(f"displayed FC does not match pair owner: {line!r}")
    return Region(
        kind=match.group("kind"),
        rg=match.group("rg"),
        pair=match.group("pair"),
        chrom=expected[0],
        start=expected[1],
        end=expected[2],
        strand=expected[3],
    )


def remap_region(
    displayed: Region,
    old_catalog: RegionCatalog,
    new_catalog: RegionCatalog,
) -> list[Region]:
    old_owner = old_catalog.by_label[f"FC:{displayed.pair}"]
    new_owners = new_catalog.fc_by_coord.get(old_owner.coord_key, [])
    if len(new_owners) != 1:
        raise ValueError(
            f"expected one refreshed FC for {old_owner.display()}, observed {new_owners}"
        )
    new_owner = new_owners[0]
    if displayed.kind == "FC":
        return [new_owner]
    matches = new_catalog.by_pair_kind_coord.get(
        (new_owner.pair, displayed.kind, displayed.coord_key), []
    )
    if not matches:
        raise ValueError(
            f"refreshed pair {new_owner.pair} lost counterpart {displayed.display()}"
        )
    return matches


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open() as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def ordered_unique_regions(regions: list[Region]) -> list[Region]:
    return list(dict.fromkeys(regions))


def ordered_unique(values: list[str]) -> list[str]:
    return list(dict.fromkeys(values))


def remap_display_value(
    value: object,
    old_catalog: RegionCatalog,
    new_catalog: RegionCatalog,
) -> list[Region]:
    remapped: list[Region] = []
    for line in str(value or "").splitlines():
        if not line or DISPLAY_RE.fullmatch(line) is None:
            continue
        displayed = parse_displayed_region(line, old_catalog)
        remapped.extend(remap_region(displayed, old_catalog, new_catalog))
    return ordered_unique_regions(remapped)


def remap_pair_value(
    value: object,
    old_catalog: RegionCatalog,
    new_catalog: RegionCatalog,
) -> list[str]:
    remapped: list[str] = []
    for pair in str(value or "").splitlines():
        if not pair:
            continue
        old_owner = old_catalog.by_label.get(f"FC:{pair}")
        if old_owner is None:
            raise ValueError(f"old pair owner FC:{pair} is absent")
        new_owners = new_catalog.fc_by_coord.get(old_owner.coord_key, [])
        if len(new_owners) != 1:
            raise ValueError(
                f"expected one refreshed FC for {old_owner.display()}, observed {new_owners}"
            )
        remapped.append(new_owners[0].pair)
    return ordered_unique(remapped)


def main() -> None:
    args = parse_args()
    new_roots = parse_assignments(args.new_pipeline_root, "--new-pipeline-root")
    haplotype_dirs = parse_assignments(args.haplotype_dir, "--haplotype-dir")
    if set(new_roots) != set(haplotype_dirs):
        raise ValueError("new pipeline roots and haplotype directories need identical assemblies")

    old_roots = dict(OLD_PIPELINE_ROOTS)
    old_roots.update(parse_assignments(args.old_pipeline_root, "--old-pipeline-root"))
    unknown = set(new_roots) - set(old_roots)
    if unknown:
        raise ValueError(f"missing old pipeline root(s): {sorted(unknown)}")

    workbook = openpyxl.load_workbook(args.workbook)
    if workbook.sheetnames != ["FN observations", "Qname mappings"]:
        raise ValueError(f"unexpected workbook sheets: {workbook.sheetnames}")
    observations = workbook["FN observations"]
    mappings = workbook["Qname mappings"]
    if [cell.value for cell in observations[1]] not in (
        OBSERVATION_COLUMNS,
        OBSERVATION_COLUMNS + [OBSERVATION_AUDIT_COLUMN],
    ):
        raise ValueError("FN observations schema differs from the canonical format")
    if [cell.value for cell in mappings[1]] not in (
        QNAME_COLUMNS,
        QNAME_COLUMNS + [QNAME_AUDIT_COLUMN],
    ):
        raise ValueError("Qname mappings schema differs from the canonical format")

    workbook_sites: dict[str, set[str]] = defaultdict(set)
    for row in observations.iter_rows(min_row=2, values_only=True):
        workbook_sites[str(row[0])].add(str(row[1]))

    refreshed = 0
    for assembly in sorted(new_roots):
        old_catalog = read_catalog(old_roots[assembly])
        new_catalog = read_catalog(new_roots[assembly])
        haplotype_dir = haplotype_dirs[assembly]

        rg_rows = read_tsv(haplotype_dir / "target_rg_summary.tsv")
        overlap_rows = read_tsv(haplotype_dir / "target_rg_overlaps.tsv")
        depth_rows = read_tsv(haplotype_dir / "pre_fp_raw_per_rg_summary.tsv")
        rgs_by_site = {
            row["FN_site"]: [rg for rg in row["relevant_RGs"].split(",") if rg]
            for row in rg_rows
        }
        if set(rgs_by_site) != workbook_sites[assembly]:
            missing = sorted(set(rgs_by_site) - workbook_sites[assembly])
            stale = sorted(workbook_sites[assembly] - set(rgs_by_site))
            raise ValueError(
                f"{assembly} FN set changed; rebuild qname evidence before refreshing "
                f"the workbook (new={missing}, stale={stale})"
            )

        target_intervals: dict[tuple[str, str], list[str]] = defaultdict(list)
        for row in overlap_rows:
            key = (row["FN_site"], row["RG"])
            if row["target_interval"] not in target_intervals[key]:
                target_intervals[key].append(row["target_interval"])

        depth_by_site_rg = {
            (row["FN_site"], row["RG"]): (
                int(row["haplotype_records_at_target"]),
                int(row["target_covering_records"]),
            )
            for row in depth_rows
        }

        for row_number in range(2, observations.max_row + 1):
            if observations.cell(row_number, 1).value != assembly:
                continue
            site = str(observations.cell(row_number, 2).value)
            relevant_rgs = sorted(rgs_by_site[site], key=rg_sort_key)

            query_regions = remap_display_value(
                observations.cell(row_number, 3).value,
                old_catalog,
                new_catalog,
            )
            counterpart_regions = remap_display_value(
                observations.cell(row_number, 4).value,
                old_catalog,
                new_catalog,
            )

            query_items = [(region.rg, region.display()) for region in query_regions]
            represented_rgs = {region.rg for region in query_regions}
            for rg in relevant_rgs:
                if rg in represented_rgs:
                    continue
                intervals = target_intervals.get((site, rg), [])
                if not intervals:
                    raise ValueError(f"{assembly} {site} has no target interval for {rg}")
                query_items.extend((rg, f"{rg} | {interval}") for interval in intervals)

            query_lines = [
                value
                for _, value in sorted(query_items, key=lambda item: rg_sort_key(item[0]))
            ]
            counterpart_regions = sorted(
                counterpart_regions, key=lambda region: rg_sort_key(region.rg)
            )

            query_pairs = remap_pair_value(
                observations.cell(row_number, 6).value,
                old_catalog,
                new_catalog,
            )
            counterpart_pairs = remap_pair_value(
                observations.cell(row_number, 7).value,
                old_catalog,
                new_catalog,
            )
            depth_lines: list[str] = []
            for rg in relevant_rgs:
                key = (site, rg)
                if key not in depth_by_site_rg:
                    raise ValueError(f"missing strict haplotype depth for {assembly} {site} {rg}")
                alt_depth, total_depth = depth_by_site_rg[key]
                depth_lines.append(f"{rg}: {alt_depth}/{total_depth}")

            observations.cell(row_number, 3, "\n".join(query_lines) or None)
            observations.cell(
                row_number,
                4,
                "\n".join(region.display() for region in counterpart_regions) or None,
            )
            observations.cell(row_number, 5, "\n".join(relevant_rgs))
            observations.cell(row_number, 6, "\n".join(query_pairs) or None)
            observations.cell(row_number, 7, "\n".join(counterpart_pairs) or None)
            observations.cell(row_number, 10, "\n".join(depth_lines))
            refreshed += 1

    print(f"workbook={args.workbook}")
    print(f"assemblies={','.join(sorted(new_roots))}")
    print(f"refreshed_FN_rows={refreshed}")
    if args.dry_run:
        print("saved=false")
        return

    output = args.output or args.workbook
    output.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = output.with_name(f".{output.name}.{os.getpid()}.tmp")
    workbook.save(tmp_path)
    os.replace(tmp_path, output)
    print(f"output={output.resolve()}")
    print("saved=true")


if __name__ == "__main__":
    main()
