#!/usr/bin/env python3
"""Add exact realigned-allele assertions and FN causes to the canonical workbook."""

from __future__ import annotations

import argparse
import copy
import csv
import os
import re
import subprocess
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path

import openpyxl


OBSERVATION_COLUMNS = [
    "Assembly",
    "FN allele",
    "Query SD region",
    "Homologous counterpart region",
    "RG index",
    "FC query pair index",
    "NFC/counterpart pair index",
    "Selected input BAM ALT-haplotype qnames",
    "Original raw BAM ALT-haplotype qnames",
    "ALT-haplotype depth / total depth at true ALT site",
]
OBSERVATION_AUDIT_COLUMN = "Realigned exact-allele AD distribution and FN cause"

QNAME_COLUMNS = [
    "Assembly",
    "FN allele",
    "ALT-haplotype qname",
    "Evidence source",
    "Selected input BAM alignment records",
    "Original raw BAM alignment records",
    "Realigned BAM alignment records",
    "Should be extracted (FC/NFC pair)",
    "Present in extracted RG FASTQ",
]
QNAME_AUDIT_COLUMN = "Realigned read pair produces exact FN ALT"

ROOT = Path("/paedyl01/disk1/yangyxt/SDrecall-rust-migration")
AUDIT_ROOT = ROOT / "test_tmp/nfc_fc_conflict_fix_20260723"
PIPELINE_ROOTS = {
    "hg38": AUDIT_ROOT / "runs/hg38/HG002_hg38_nfcfix23_SDrecall",
    "hg19": AUDIT_ROOT / "runs/hg19/HG002_hg19_nfcfix23_SDrecall",
    "chm13": AUDIT_ROOT / "runs/t2t/HG002_chm13_nfcfix23_SDrecall",
}
HAPLOTYPE_DIRS = {
    assembly: AUDIT_ROOT / "haplotypes" / assembly for assembly in PIPELINE_ROOTS
}
REFERENCES = {
    "hg38": Path("/paedyl01/disk1/yangyxt/indexed_genome/hg38/ucsc.hg38.fasta"),
    "hg19": Path("/paedyl01/disk1/yangyxt/indexed_genome/ucsc.hg19.fasta"),
    "chm13": Path(
        "/paedyl01/disk1/yangyxt/indexed_genome/chm13/chm13.draft_v1.1.fasta"
    ),
}

QNAME_RG_SUFFIX_RE = re.compile(r":RG\d+$")
QNAME_MATE_SUFFIX_RE = re.compile(r"/[12]$")
RG_RE = re.compile(r"^RG\d+$")


@dataclass(frozen=True)
class Site:
    assembly: str
    chrom: str
    pos: int
    ref: str
    alt: str

    @property
    def label(self) -> str:
        return f"{self.chrom}:{self.pos}:{self.ref}:{self.alt}"


@dataclass(frozen=True)
class PileupEvent:
    base: str
    indels: tuple[tuple[str, str], ...]


@dataclass(frozen=True)
class AlleleDepth:
    rg: str
    depth: int
    ref: int
    alt: int
    other: int
    alt_qnames: frozenset[str]

    @property
    def alt_fraction(self) -> float:
        return self.alt / self.depth if self.depth else 0.0


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--workbook",
        type=Path,
        default=ROOT / "FN_CAUSE_DIAGNOSIS.xlsx",
    )
    parser.add_argument("--output", type=Path)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--samtools", default="samtools")
    parser.add_argument(
        "--pipeline-root",
        action="append",
        default=[],
        metavar="ASSEMBLY=PATH",
    )
    parser.add_argument(
        "--haplotype-dir",
        action="append",
        default=[],
        metavar="ASSEMBLY=PATH",
    )
    parser.add_argument(
        "--reference",
        action="append",
        default=[],
        metavar="ASSEMBLY=PATH",
    )
    return parser.parse_args()


def apply_path_overrides(
    paths: dict[str, Path], assignments: list[str], option: str
) -> None:
    for assignment in assignments:
        assembly, separator, value = assignment.partition("=")
        if not separator or assembly not in paths or not value:
            raise ValueError(
                f"{option} expects one of "
                f"{', '.join(sorted(paths))}=PATH; observed {assignment!r}"
            )
        paths[assembly] = Path(value)


def worksheet_headers(ws: openpyxl.worksheet.worksheet.Worksheet) -> list[str]:
    return [cell.value for cell in ws[1]]


def canonical_qname(name: str) -> str:
    name = name.strip().split(None, 1)[0]
    name = QNAME_MATE_SUFFIX_RE.sub("", name)
    return QNAME_RG_SUFFIX_RE.sub("", name)


def parse_site(assembly: str, value: object) -> Site:
    fields = str(value).split(":", 3)
    if len(fields) != 4:
        raise ValueError(f"unexpected FN allele label: {value!r}")
    chrom, pos, ref, alt = fields
    site = Site(assembly, chrom, int(pos), ref.upper(), alt.upper())
    if len(site.ref) == len(site.alt) and len(site.ref) != 1:
        raise ValueError(f"multi-base substitutions are not supported: {site.label}")
    is_insertion = len(site.alt) > len(site.ref) and site.alt.startswith(site.ref)
    is_deletion = len(site.ref) > len(site.alt) and site.ref.startswith(site.alt)
    if len(site.ref) != len(site.alt) and not (is_insertion or is_deletion):
        raise ValueError(f"non-normalized or complex FN allele: {site.label}")
    return site


def parse_rgs(value: object) -> list[str]:
    values = re.split(r"[;\s]+", str(value or ""))
    return sorted(
        {value for value in values if RG_RE.fullmatch(value)},
        key=lambda rg: int(rg[2:]),
    )


def parse_pileup_events(bases: str) -> list[PileupEvent]:
    events: list[PileupEvent] = []
    index = 0
    while index < len(bases):
        if bases[index] == "^":
            index += 2
            continue
        if bases[index] == "$":
            index += 1
            continue

        base = bases[index]
        index += 1
        indels: list[tuple[str, str]] = []
        while index < len(bases) and bases[index] in "+-":
            operation = bases[index]
            index += 1
            length_start = index
            while index < len(bases) and bases[index].isdigit():
                index += 1
            if index == length_start:
                raise ValueError(f"malformed mpileup indel in {bases!r}")
            length = int(bases[length_start:index])
            sequence = bases[index : index + length]
            if len(sequence) != length:
                raise ValueError(f"truncated mpileup indel in {bases!r}")
            index += length
            indels.append((operation, sequence.upper()))
        events.append(PileupEvent(base, tuple(indels)))
    return events


def event_call(event: PileupEvent, site: Site) -> str:
    base = event.base.upper()
    anchor_is_ref = event.base in ".," or base == site.ref[0]
    if len(site.ref) == len(site.alt):
        if base == site.alt:
            return "alt"
        if anchor_is_ref:
            return "ref"
        return "other"

    if len(site.alt) > len(site.ref):
        expected = ("+", site.alt[len(site.ref) :])
    else:
        expected = ("-", site.ref[len(site.alt) :])
    if anchor_is_ref and expected in event.indels:
        return "alt"
    if anchor_is_ref and not event.indels:
        return "ref"
    return "other"


def run_mpileup(
    samtools: str,
    bam: Path,
    reference: Path,
    site: Site,
    rg: str,
) -> AlleleDepth:
    region = f"{site.chrom}:{site.pos}-{site.pos}"
    command = [
        samtools,
        "mpileup",
        "-A",
        "-q",
        "10",
        "-Q",
        "15",
        "--output-QNAME",
        "-f",
        str(reference),
        "-r",
        region,
        str(bam),
    ]
    completed = subprocess.run(
        command,
        check=True,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    output_lines = [line for line in completed.stdout.splitlines() if line]
    if not output_lines:
        return AlleleDepth(rg, 0, 0, 0, 0, frozenset())
    if len(output_lines) != 1:
        raise ValueError(f"expected one mpileup row for {site.label} {rg}")

    fields = output_lines[0].split("\t")
    if len(fields) < 7 or fields[0] != site.chrom or int(fields[1]) != site.pos:
        raise ValueError(f"unexpected mpileup output for {site.label} {rg}")
    reported_depth = int(fields[3])
    events = parse_pileup_events(fields[4])
    qnames = [] if fields[6] == "*" else fields[6].split(",")
    if reported_depth != len(events) or len(events) != len(qnames):
        raise ValueError(
            f"mpileup depth/event/qname mismatch for {site.label} {rg}: "
            f"{reported_depth}/{len(events)}/{len(qnames)}"
        )

    counts: Counter[str] = Counter()
    alt_qnames: set[str] = set()
    for event, qname in zip(events, qnames):
        call = event_call(event, site)
        counts[call] += 1
        if call == "alt":
            alt_qnames.add(canonical_qname(qname))
    return AlleleDepth(
        rg,
        reported_depth,
        counts["ref"],
        counts["alt"],
        counts["other"],
        frozenset(alt_qnames),
    )


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open() as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def copy_cell_style(source: openpyxl.cell.cell.Cell, target: openpyxl.cell.cell.Cell) -> None:
    if source.has_style:
        target._style = copy.copy(source._style)
    target.font = copy.copy(source.font)
    target.fill = copy.copy(source.fill)
    target.border = copy.copy(source.border)
    target.alignment = copy.copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy.copy(source.protection)


def qname_assertion(
    row: tuple[object, ...],
    exact_rgs: set[str],
) -> str:
    should_extract = str(row[7] or "").startswith("Yes |")
    in_fastq = str(row[8] or "").startswith("Yes |")
    in_realigned_bam = bool(row[6])
    if not should_extract:
        return "Not extracted | not eligible"
    if not in_fastq:
        return "No | eligible pair absent from extracted RG FASTQ"
    if exact_rgs:
        return f"Yes | {','.join(sorted(exact_rgs, key=lambda rg: int(rg[2:])))}"
    if not in_realigned_bam:
        return "No | extracted pair absent from realigned BAM"
    return "No | realigned pair does not encode exact ALT at target"


def classify_cause(
    mapping_rows: list[tuple[object, ...]],
    input_marker_qnames: int,
    strict_haplotype_records: int,
    exact_alt_records: int,
    audited_exact_qnames: int,
    raw_vcf_has_allele: bool,
) -> tuple[str, str]:
    eligible = [row for row in mapping_rows if str(row[7] or "").startswith("Yes |")]
    in_fastq = [row for row in eligible if str(row[8] or "").startswith("Yes |")]
    in_realigned_bam = [row for row in in_fastq if row[6]]

    if input_marker_qnames == 0:
        return (
            "Cat1",
            "selected input BAM has no verified full ALT haplotype; loss precedes extraction",
        )
    if not eligible:
        overlaps_network = any(
            re.search(r"\b(?:FC|NFC) overlapped(?:;|$)", str(row[7] or ""))
            is not None
            for row in mapping_rows
            if str(row[3]) == "input+raw"
        )
        if overlaps_network:
            return (
                "Cat3",
                "ALT qname reaches a relevant FC/NFC interval, but no complete eligible read pair is formed",
            )
        return (
            "Cat2",
            "selected-input ALT haplotype exists, but no ALT qname is eligible in the relevant FC/NFC network",
        )
    if not in_fastq:
        return "Cat3", "eligible ALT read pair is absent from the extracted RG FASTQ"
    if not in_realigned_bam:
        return "Cat3", "extracted ALT read pair is absent from the realigned RG BAM"
    if strict_haplotype_records == 0:
        if exact_alt_records:
            return (
                "Cat4",
                "extracted reads are realigned, but no verified full ALT haplotype reaches the strict target; exact ALT observations are isolated evidence only",
            )
        return (
            "Cat4",
            "extracted reads are realigned, but neither the verified full ALT haplotype nor the exact ALT survives at the target",
        )
    if audited_exact_qnames == 0:
        if exact_alt_records:
            return (
                "Cat4",
                "verified full ALT sequence reaches the target, but no verified ALT-haplotype qname encodes the exact normalized ALT under caller thresholds; other exact observations are unlinked",
            )
        return (
            "Cat4",
            "verified full ALT sequence reaches the target, but no alignment encodes the exact normalized ALT under caller thresholds",
        )
    if not raw_vcf_has_allele:
        return (
            "Cat5",
            "verified full ALT haplotype and exact ALT survive pre-FP realignment, but the raw production VCF does not emit the allele",
        )
    return (
        "Downstream",
        "raw production VCF emits the exact ALT, but the benchmark callset loses it downstream",
    )


def format_observation(
    depths: list[AlleleDepth],
    category: str,
    cause: str,
) -> str:
    lines = [
        (
            f"{depth.rg}: ALT={depth.alt}, REF={depth.ref}, other={depth.other}, "
            f"DP={depth.depth}, AF={depth.alt_fraction:.2%}"
        )
        for depth in depths
    ]
    lines.append(f"Cause: {category} | {cause}")
    return "\n".join(lines)


def main() -> None:
    args = parse_args()
    apply_path_overrides(PIPELINE_ROOTS, args.pipeline_root, "--pipeline-root")
    apply_path_overrides(HAPLOTYPE_DIRS, args.haplotype_dir, "--haplotype-dir")
    apply_path_overrides(REFERENCES, args.reference, "--reference")

    workbook = openpyxl.load_workbook(args.workbook)
    if workbook.sheetnames != ["FN observations", "Qname mappings"]:
        raise ValueError(f"unexpected workbook sheets: {workbook.sheetnames}")
    observations = workbook["FN observations"]
    mappings = workbook["Qname mappings"]
    if worksheet_headers(observations) not in (
        OBSERVATION_COLUMNS,
        OBSERVATION_COLUMNS + [OBSERVATION_AUDIT_COLUMN],
    ):
        raise ValueError("FN observations schema differs from the agreed format")
    if worksheet_headers(mappings) not in (
        QNAME_COLUMNS,
        QNAME_COLUMNS + [QNAME_AUDIT_COLUMN],
    ):
        raise ValueError("Qname mappings schema differs from the agreed format")

    observation_rows = list(
        observations.iter_rows(min_row=2, max_col=len(OBSERVATION_COLUMNS), values_only=True)
    )
    mapping_rows = list(
        mappings.iter_rows(min_row=2, max_col=len(QNAME_COLUMNS), values_only=True)
    )
    mappings_by_site: dict[tuple[str, str], list[tuple[object, ...]]] = defaultdict(list)
    for row in mapping_rows:
        mappings_by_site[(str(row[0]), str(row[1]))].append(row)

    pre_by_site: dict[tuple[str, str], dict[str, str]] = {}
    strict_by_site_rg: dict[tuple[str, str, str], dict[str, str]] = {}
    raw_vcf_by_site: dict[tuple[str, str], bool] = {}
    for assembly, haplotype_dir in HAPLOTYPE_DIRS.items():
        for row in read_tsv(haplotype_dir / "pre_fp_input_qname_summary.tsv"):
            pre_by_site[(assembly, row["FN_site"])] = row
        for row in read_tsv(haplotype_dir / "pre_fp_raw_per_rg_summary.tsv"):
            strict_by_site_rg[(assembly, row["FN_site"], row["RG"])] = row
        for row in read_tsv(haplotype_dir / "vcf_stage_exact_allele_presence.tsv"):
            if row["stage"] == "raw":
                raw_vcf_by_site[(assembly, row["FN_site"])] = (
                    row["exact_allele_present"] == "yes"
                )

    exact_rgs_by_qname: dict[tuple[str, str, str], set[str]] = defaultdict(set)
    calculated_observations: list[str] = []
    categories: Counter[tuple[str, str]] = Counter()
    for row in observation_rows:
        assembly = str(row[0])
        site = parse_site(assembly, row[1])
        key = (assembly, site.label)
        rgs = parse_rgs(row[4])
        if not rgs:
            raise ValueError(f"no relevant RGs in workbook for {assembly} {site.label}")

        depths: list[AlleleDepth] = []
        for rg in rgs:
            bam = (
                PIPELINE_ROOTS[assembly]
                / "recall_results"
                / f"HG002.sdrecall.only_{rg}.raw.bam"
            )
            if not bam.is_file() or not Path(f"{bam}.bai").is_file():
                raise FileNotFoundError(f"missing indexed realigned BAM: {bam}")
            depth = run_mpileup(args.samtools, bam, REFERENCES[assembly], site, rg)
            depths.append(depth)
            for qname in depth.alt_qnames:
                exact_rgs_by_qname[(assembly, site.label, qname)].add(rg)

        pre = pre_by_site.get(key)
        if pre is None:
            raise ValueError(f"missing pre-FP qname summary for {assembly} {site.label}")
        strict_rows = [strict_by_site_rg.get((assembly, site.label, rg)) for rg in rgs]
        if any(strict_row is None for strict_row in strict_rows):
            raise ValueError(f"missing strict per-RG summary for {assembly} {site.label}")
        strict_records = sum(
            int(strict_row["haplotype_records_at_target"])
            for strict_row in strict_rows
            if strict_row is not None
        )
        exact_alt_records = sum(depth.alt for depth in depths)
        exact_qnames = set().union(*(depth.alt_qnames for depth in depths))
        audited_exact_qnames = sum(
            canonical_qname(str(mapping_row[2])) in exact_qnames
            and str(mapping_row[7] or "").startswith("Yes |")
            and str(mapping_row[8] or "").startswith("Yes |")
            for mapping_row in mappings_by_site[key]
        )
        category, cause = classify_cause(
            mappings_by_site[key],
            int(pre["input_marker_qnames"]),
            strict_records,
            exact_alt_records,
            audited_exact_qnames,
            raw_vcf_by_site[key],
        )
        categories[(assembly, category)] += 1
        calculated_observations.append(format_observation(depths, category, cause))
        depth_text = "; ".join(
            f"{depth.rg}={depth.alt}/{depth.depth} ({depth.alt_fraction:.2%})"
            for depth in depths
        )
        print(
            f"ALLELE\t{assembly}\t{site.label}\t{category}\t{depth_text}\t"
            f"audited_exact_qnames={audited_exact_qnames}\t{cause}"
        )

    calculated_qnames: list[str] = []
    for row in mapping_rows:
        key = (str(row[0]), str(row[1]), canonical_qname(str(row[2])))
        calculated_qnames.append(qname_assertion(row, exact_rgs_by_qname.get(key, set())))

    print(f"workbook={args.workbook.resolve()}")
    print(f"FN_observation_rows={len(observation_rows)}")
    print(f"qname_mapping_rows={len(mapping_rows)}")
    print(f"qname_exact_ALT_yes_rows={sum(value.startswith('Yes |') for value in calculated_qnames)}")
    for (assembly, category), count in sorted(categories.items()):
        print(f"CATEGORY\t{assembly}\t{category}\t{count}")
    if len(observation_rows) != 45 or len(mapping_rows) != 3632:
        raise ValueError("canonical workbook row counts changed")
    if args.dry_run:
        print("saved=false")
        return

    observation_column = len(OBSERVATION_COLUMNS) + 1
    qname_column = len(QNAME_COLUMNS) + 1
    observation_header = observations.cell(
        1, observation_column, OBSERVATION_AUDIT_COLUMN
    )
    copy_cell_style(observations.cell(1, observation_column - 1), observation_header)
    for row_number, value in enumerate(calculated_observations, start=2):
        cell = observations.cell(row_number, observation_column, value)
        copy_cell_style(observations.cell(row_number, observation_column - 1), cell)

    qname_header = mappings.cell(1, qname_column, QNAME_AUDIT_COLUMN)
    copy_cell_style(mappings.cell(1, qname_column - 1), qname_header)
    for row_number, value in enumerate(calculated_qnames, start=2):
        cell = mappings.cell(row_number, qname_column, value)
        copy_cell_style(mappings.cell(row_number, qname_column - 1), cell)

    observations.column_dimensions["K"].width = 88
    mappings.column_dimensions["J"].width = 48
    observations.auto_filter.ref = f"A1:K{observations.max_row}"
    mappings.auto_filter.ref = f"A1:J{mappings.max_row}"

    output = args.output or args.workbook
    output.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = output.with_name(f".{output.name}.{os.getpid()}.tmp")
    workbook.save(tmp_path)
    os.replace(tmp_path, output)
    print(f"output={output.resolve()}")
    print("saved=true")


if __name__ == "__main__":
    main()
