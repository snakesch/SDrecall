#!/usr/bin/env python3
"""Trace every eligible Cat4 ALT-haplotype qname from extraction to realignment.

The production FASTQ stores a qname but not the FC/NFC interval that caused it
to be included.  Consequently, this audit reports every eligible extraction
path and explicitly marks qnames with more than one possible source path.
"""

from __future__ import annotations

import argparse
import csv
import os
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import openpyxl
import pysam

from update_fn_realign_allele_audit import (
    REFERENCES,
    Site,
    canonical_qname,
    parse_site,
    run_mpileup,
)


ROOT = Path("/paedyl01/disk1/yangyxt/SDrecall-rust-migration")
AUDIT_ROOT = ROOT / "test_tmp/nfc_fc_conflict_fix_20260723"
PIPELINE_ROOTS = {
    "hg38": AUDIT_ROOT / "runs/hg38/HG002_hg38_nfcfix23_SDrecall",
    "hg19": AUDIT_ROOT / "runs/hg19/HG002_hg19_nfcfix23_SDrecall",
    "chm13": AUDIT_ROOT / "runs/t2t/HG002_chm13_nfcfix23_SDrecall",
}
HAPLOTYPE_DIRS = {
    assembly: AUDIT_ROOT / "haplotypes" / assembly for assembly in PIPELINE_ROOTS
}

OBSERVATION_COLUMNS = [
    "Assembly",
    "FN allele",
    "Query SD region",
    "Homologous counterpart region",
    "RG index",
    "FC query pair index",
    "NFC/counterpart pair index",
    "Selected input BAM ALT-haplotype qnames",
    "Original raw BAM ALT-haplotype qnames",
    "ALT-haplotype depth / total depth at true ALT site",
    "Realigned exact-allele AD distribution and FN cause",
]
QNAME_COLUMNS = [
    "Assembly",
    "FN allele",
    "ALT-haplotype qname",
    "Evidence source",
    "Selected input BAM alignment records",
    "Original raw BAM alignment records",
    "Realigned BAM alignment records",
    "Should be extracted (FC/NFC pair)",
    "Present in extracted RG FASTQ",
    "Realigned read pair produces exact FN ALT",
]

SOURCE_RE = re.compile(r"(?:^|; )Yes \| (?P<labels>(?:FC|NFC):RG\d+_\d+(?:,(?:FC|NFC):RG\d+_\d+)*)")
LABEL_RE = re.compile(r"^(?P<kind>FC|NFC):(?P<pair>RG\d+_\d+)$")
RG_RE = re.compile(r"^RG\d+$")


@dataclass(frozen=True)
class Interval:
    chrom: str
    start: int
    end: int

    def overlaps(self, other: "Interval") -> bool:
        return self.chrom == other.chrom and self.start < other.end and other.start < self.end

    def contains(self, chrom: str, position0: int) -> bool:
        return self.chrom == chrom and self.start <= position0 < self.end

    def display(self) -> str:
        return f"{self.chrom}:{self.start + 1}-{self.end}"


@dataclass(frozen=True)
class SourcePath:
    kind: str
    pair: str

    @property
    def rg(self) -> str:
        return self.pair.split("_", 1)[0]

    @property
    def label(self) -> str:
        return f"{self.kind}:{self.pair}"

    @property
    def intended_fc(self) -> str:
        return f"FC:{self.pair}"


@dataclass(frozen=True)
class TraceRecord:
    rg: str
    qname: str
    record_id: str
    interval: Interval
    mapq: int
    cigar: str
    marker_bearing: bool
    overlaps_target: bool
    strict_marker_at_target: bool

    def display(self) -> str:
        marker = "marker" if self.marker_bearing else "mate"
        return (
            f"{self.rg}:{self.record_id.rsplit('|', 1)[-1]}:{self.interval.display()}:"
            f"MQ{self.mapq}:{self.cigar}:{marker}"
        )


@dataclass(frozen=True)
class ExactRecord:
    rg: str
    record_id: str
    mapq: int
    cigar: str
    minimum_event_bq: int | None

    def display(self) -> str:
        bq = "." if self.minimum_event_bq is None else str(self.minimum_event_bq)
        return f"{self.rg}:{self.record_id.rsplit('|', 1)[-1]}:MQ{self.mapq}:BQ{bq}:{self.cigar}"


@dataclass
class QnameAudit:
    assembly: str
    site: Site
    qname: str
    evidence_source: str
    sources: list[SourcePath]
    fastq_status: str
    trace_records: list[TraceRecord]
    exact_records: list[ExactRecord]
    production_exact_rgs: set[str]
    fc_catalog: dict[str, list[Interval]]

    @property
    def marker_records(self) -> list[TraceRecord]:
        return [record for record in self.trace_records if record.marker_bearing]

    @property
    def source_provenance(self) -> str:
        if len(self.sources) == 1:
            return "unambiguous"
        return f"ambiguous ({len(self.sources)} eligible paths; FASTQ has no source tag)"

    def fc_overlaps(self, records: Iterable[TraceRecord]) -> set[str]:
        labels: set[str] = set()
        for record in records:
            for label, intervals in self.fc_catalog.items():
                if label.startswith(f"FC:{record.rg}_") and any(
                    record.interval.overlaps(interval) for interval in intervals
                ):
                    labels.add(label)
        return labels

    @property
    def marker_fc_overlaps(self) -> set[str]:
        return self.fc_overlaps(self.marker_records)

    @property
    def pair_fc_overlaps(self) -> set[str]:
        return self.fc_overlaps(self.trace_records)

    @property
    def marker_overlaps_fn(self) -> bool:
        return any(record.overlaps_target for record in self.marker_records)

    @property
    def exact_record_ids(self) -> set[str]:
        return {record.record_id for record in self.exact_records}

    @property
    def marker_record_ids(self) -> set[str]:
        return {record.record_id for record in self.marker_records}

    @property
    def marker_linked_exact_records(self) -> list[ExactRecord]:
        marker_ids = self.marker_record_ids
        return [record for record in self.exact_records if record.record_id in marker_ids]

    @property
    def marker_linked_production_exact(self) -> bool:
        return any(
            record.rg in self.production_exact_rgs
            for record in self.marker_linked_exact_records
        )

    def source_contains_fn(self, source: SourcePath) -> bool:
        return any(
            interval.contains(self.site.chrom, self.site.pos - 1)
            for interval in self.fc_catalog.get(source.intended_fc, [])
        )

    @property
    def query_linked_sources(self) -> list[SourcePath]:
        return [source for source in self.sources if self.source_contains_fn(source)]

    def path_outcome(self, source: SourcePath) -> tuple[str, set[str]]:
        marker_fc = {
            label for label in self.marker_fc_overlaps if label.startswith(f"FC:{source.rg}_")
        }
        if source.intended_fc in marker_fc:
            others = marker_fc - {source.intended_fc}
            if others:
                return "intended FC reached; overlapping other FC(s)", others
            return "intended FC reached", set()
        if marker_fc:
            return "other FC(s) only", marker_fc
        if any(record.rg == source.rg for record in self.marker_records):
            return "outside every FC in this RG", set()
        return "ALT-haplotype record absent in this RG", set()

    def placement_summary(self) -> str:
        path_text: list[str] = []
        for source in self.sources:
            outcome, others = self.path_outcome(source)
            suffix = f" [{','.join(sorted(others))}]" if others else ""
            relation = "query-linked" if self.source_contains_fn(source) else "incidental"
            path_text.append(f"{source.label}[{relation}]->{outcome}{suffix}")
        return "; ".join(path_text)

    def assertion(self) -> str:
        if self.marker_linked_production_exact:
            exact = "Yes | linked ALT-haplotype record passes MAPQ>=10/BQ>=15"
        elif self.marker_linked_exact_records:
            exact = (
                "No | exact ALT reaches target but fails MAPQ>=10/BQ>=15: "
                + ",".join(record.display() for record in self.marker_linked_exact_records)
            )
        elif self.marker_overlaps_fn:
            exact = "No | ALT-haplotype record reaches FN coordinate but does not encode exact ALT"
        elif self.marker_records:
            exact = "No | ALT-haplotype record does not reach FN coordinate"
        else:
            exact = "No | realigned pair has no record retaining the verified ALT haplotype"
        return (
            f"{exact}\nPath provenance: {self.source_provenance}\n"
            f"FN-linked extraction path: {'yes' if self.query_linked_sources else 'no'}\n"
            f"Placement: {self.placement_summary()}"
        )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=ROOT / "FN_CAUSE_DIAGNOSIS.xlsx")
    parser.add_argument("--output-workbook", type=Path)
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=AUDIT_ROOT / "cat4_qname_audit",
    )
    parser.add_argument("--update-workbook", action="store_true")
    parser.add_argument("--samtools", default="samtools")
    return parser.parse_args()


def worksheet_headers(ws: openpyxl.worksheet.worksheet.Worksheet) -> list[str]:
    return [cell.value for cell in ws[1]]


def parse_source_paths(value: object) -> list[SourcePath]:
    paths: list[SourcePath] = []
    seen: set[str] = set()
    for match in SOURCE_RE.finditer(str(value or "")):
        for label in match.group("labels").split(","):
            label_match = LABEL_RE.fullmatch(label)
            if label_match is None or label in seen:
                continue
            seen.add(label)
            paths.append(SourcePath(label_match.group("kind"), label_match.group("pair")))
    return sorted(paths, key=lambda path: (int(path.rg[2:]), path.kind, path.pair))


def read_fc_catalog(run_root: Path, rgs: set[str]) -> dict[str, list[Interval]]:
    catalog: dict[str, list[Interval]] = defaultdict(list)
    for rg in sorted(rgs, key=lambda value: int(value[2:])):
        path = run_root / "realign_groups" / rg / f"{rg}_related_homo_regions.bed"
        if not path.is_file():
            raise FileNotFoundError(path)
        with path.open() as handle:
            for line in handle:
                fields = line.rstrip("\n").split("\t")
                if len(fields) < 7 or not fields[6].startswith("FC:"):
                    continue
                catalog[fields[6]].append(
                    Interval(fields[0], int(fields[1]), int(fields[2]))
                )
    return dict(catalog)


def read_trace_rows(
    assembly: str,
    path: Path,
    cat4_sites: set[str],
) -> dict[tuple[str, str, str], list[TraceRecord]]:
    trace: dict[tuple[str, str, str], list[TraceRecord]] = defaultdict(list)
    with path.open() as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        for row in reader:
            site = row["FN_site"]
            if site not in cat4_sites:
                continue
            qname = canonical_qname(row["read_qname"])
            record = TraceRecord(
                rg=row["RG"],
                qname=qname,
                record_id=row["read_record_id"],
                interval=Interval(
                    row["mapping_chrom"],
                    int(row["mapping_start"]) - 1,
                    int(row["mapping_end"]),
                ),
                mapq=int(row["MAPQ"]),
                cigar=row["CIGAR"],
                marker_bearing=row["exact_marker_in_record"] == "yes",
                overlaps_target=row["marker_alignment_overlaps_target"] == "yes",
                strict_marker_at_target=row["strict_marker_at_target"] == "yes",
            )
            trace[(assembly, site, qname)].append(record)
    return trace


def resolve_contig(bam: pysam.AlignmentFile, chrom: str) -> str:
    if chrom in bam.references:
        return chrom
    aliases = [chrom.removeprefix("chr")]
    if not chrom.startswith("chr"):
        aliases.append(f"chr{chrom}")
    for alias in aliases:
        if alias in bam.references:
            return alias
    raise ValueError(f"contig {chrom!r} is absent from BAM")


def record_id(read: pysam.AlignedSegment) -> str:
    if read.is_read1:
        mate = "1"
    elif read.is_read2:
        mate = "2"
    else:
        mate = "0"
    return f"{canonical_qname(read.query_name)}|{mate}"


def exact_alt_event(
    read: pysam.AlignedSegment, site: Site
) -> tuple[bool, int | None]:
    if read.is_unmapped or read.query_sequence is None:
        return False, None
    sequence = read.query_sequence.upper()
    qualities = read.query_qualities
    ref_cursor = read.reference_start
    query_cursor = 0
    anchor0 = site.pos - 1
    anchor_base: str | None = None
    anchor_bq: int | None = None

    for operation, length in read.cigartuples or []:
        if operation in (0, 7, 8):  # M, =, X
            if ref_cursor <= anchor0 < ref_cursor + length:
                query_index = query_cursor + anchor0 - ref_cursor
                anchor_base = sequence[query_index]
                if qualities is not None:
                    anchor_bq = int(qualities[query_index])
            ref_cursor += length
            query_cursor += length
        elif operation == 1:  # insertion after ref_cursor - 1
            if len(site.alt) > len(site.ref) and ref_cursor - 1 == anchor0:
                inserted = sequence[query_cursor : query_cursor + length]
                expected = site.alt[len(site.ref) :]
                if anchor_base == site.alt[0] and inserted == expected:
                    event_qualities = [anchor_bq] if anchor_bq is not None else []
                    if qualities is not None:
                        event_qualities.extend(
                            int(value)
                            for value in qualities[query_cursor : query_cursor + length]
                        )
                    return True, min(event_qualities) if event_qualities else None
            query_cursor += length
        elif operation in (2, 3):  # deletion or reference skip
            if (
                operation == 2
                and len(site.ref) > len(site.alt)
                and ref_cursor == anchor0 + len(site.alt)
                and length == len(site.ref) - len(site.alt)
                and anchor_base == site.alt[0]
            ):
                return True, anchor_bq
            ref_cursor += length
        elif operation == 4:  # soft clip
            query_cursor += length
        elif operation in (5, 6):  # hard clip or padding
            continue

    if len(site.ref) == len(site.alt) == 1 and anchor_base == site.alt:
        return True, anchor_bq
    return False, anchor_bq


def target_exact_records(
    assembly: str,
    site: Site,
    rgs: set[str],
    qnames: set[str],
    samtools: str,
) -> tuple[dict[str, list[ExactRecord]], dict[str, set[str]]]:
    exact_by_qname: dict[str, list[ExactRecord]] = defaultdict(list)
    production_rgs_by_qname: dict[str, set[str]] = defaultdict(set)
    run_root = PIPELINE_ROOTS[assembly]
    for rg in sorted(rgs, key=lambda value: int(value[2:])):
        bam_path = run_root / "recall_results" / f"HG002.sdrecall.only_{rg}.raw.bam"
        with pysam.AlignmentFile(bam_path, "rb") as bam:
            contig = resolve_contig(bam, site.chrom)
            for read in bam.fetch(contig, site.pos - 1, site.pos):
                qname = canonical_qname(read.query_name or "")
                if qname not in qnames or read.is_secondary or read.is_supplementary:
                    continue
                supports, minimum_bq = exact_alt_event(read, site)
                if supports:
                    exact_by_qname[qname].append(
                        ExactRecord(
                            rg,
                            record_id(read),
                            read.mapping_quality,
                            read.cigarstring or "*",
                            minimum_bq,
                        )
                    )

        depth = run_mpileup(samtools, bam_path, REFERENCES[assembly], site, rg)
        for qname in depth.alt_qnames.intersection(qnames):
            production_rgs_by_qname[qname].add(rg)
    return exact_by_qname, production_rgs_by_qname


def fn_containing_fcs(
    site: Site, fc_catalog: dict[str, list[Interval]], rg: str
) -> set[str]:
    return {
        label
        for label, intervals in fc_catalog.items()
        if label.startswith(f"FC:{rg}_")
        and any(interval.contains(site.chrom, site.pos - 1) for interval in intervals)
    }


def path_row(audit: QnameAudit, source: SourcePath) -> dict[str, object]:
    outcome, other_fcs = audit.path_outcome(source)
    intended_intervals = audit.fc_catalog.get(source.intended_fc, [])
    target_fcs = fn_containing_fcs(audit.site, audit.fc_catalog, source.rg)
    marker_records = [record for record in audit.marker_records if record.rg == source.rg]
    pair_records = [record for record in audit.trace_records if record.rg == source.rg]
    observed_marker_fcs = {
        label
        for label in audit.marker_fc_overlaps
        if label.startswith(f"FC:{source.rg}_")
    }
    return {
        "Assembly": audit.assembly,
        "FN allele": audit.site.label,
        "ALT-haplotype qname": audit.qname,
        "Evidence source": audit.evidence_source,
        "Path provenance": audit.source_provenance,
        "Extraction source": source.label,
        "RG": source.rg,
        "Intended FC": source.intended_fc,
        "Intended FC interval": ";".join(
            interval.display() for interval in intended_intervals
        ),
        "Intended FC contains FN anchor": (
            "yes" if source.intended_fc in target_fcs else "no"
        ),
        "All FN-containing FCs in RG": ";".join(sorted(target_fcs)),
        "Present in extracted RG FASTQ": audit.fastq_status,
        "Realigned pair records": ";".join(record.display() for record in pair_records),
        "Realigned ALT-haplotype records": ";".join(
            record.display() for record in marker_records
        ),
        "FCs overlapped by ALT-haplotype record": ";".join(
            sorted(observed_marker_fcs)
        ),
        "Intended FC reached": "yes" if source.intended_fc in observed_marker_fcs else "no",
        "Other FC reached without intended FC": (
            "yes" if outcome == "other FC(s) only" else "no"
        ),
        "Other FC labels": ";".join(sorted(other_fcs)),
        "ALT-haplotype record overlaps FN anchor": (
            "yes" if any(record.overlaps_target for record in marker_records) else "no"
        ),
        "Linked exact ALT record at target": ";".join(
            record.display()
            for record in audit.marker_linked_exact_records
            if record.rg == source.rg
        ),
        "Linked exact ALT passes MAPQ10/BQ15": (
            "yes" if source.rg in audit.production_exact_rgs else "no"
        ),
        "Path outcome": outcome,
    }


def qname_row(audit: QnameAudit) -> dict[str, object]:
    outcomes = [audit.path_outcome(source)[0] for source in audit.sources]
    return {
        "Assembly": audit.assembly,
        "FN allele": audit.site.label,
        "ALT-haplotype qname": audit.qname,
        "Evidence source": audit.evidence_source,
        "Eligible extraction paths": ";".join(source.label for source in audit.sources),
        "Path provenance": audit.source_provenance,
        "Query-linked extraction path": (
            "yes" if audit.query_linked_sources else "no"
        ),
        "NFC-eligible": "yes" if any(source.kind == "NFC" for source in audit.sources) else "no",
        "FC-eligible": "yes" if any(source.kind == "FC" for source in audit.sources) else "no",
        "Present in extracted RG FASTQ": audit.fastq_status,
        "Realigned pair records": ";".join(
            record.display() for record in audit.trace_records
        ),
        "Realigned ALT-haplotype records": ";".join(
            record.display() for record in audit.marker_records
        ),
        "ALT-haplotype record FC overlaps": ";".join(sorted(audit.marker_fc_overlaps)),
        "Any intended FC reached": (
            "yes" if any(outcome.startswith("intended FC reached") for outcome in outcomes) else "no"
        ),
        "Any other-FC-only path": "yes" if "other FC(s) only" in outcomes else "no",
        "ALT-haplotype record overlaps FN anchor": "yes" if audit.marker_overlaps_fn else "no",
        "Linked exact ALT record at target": ";".join(
            record.display() for record in audit.marker_linked_exact_records
        ),
        "Linked exact ALT passes MAPQ10/BQ15": (
            "yes" if audit.marker_linked_production_exact else "no"
        ),
        "Workbook assertion": audit.assertion(),
    }


def allele_summary(
    assembly: str,
    site: Site,
    audits: list[QnameAudit],
) -> dict[str, object]:
    nfc_audits = [audit for audit in audits if any(source.kind == "NFC" for source in audit.sources)]
    source_paths = [source for audit in audits for source in audit.sources]
    nfc_paths = [source for source in source_paths if source.kind == "NFC"]
    query_nfc_paths = 0
    query_nfc_intended = 0
    query_nfc_other_only = 0
    incidental_nfc_other_only = 0
    other_transitions: Counter[str] = Counter()
    for audit in audits:
        for source in audit.sources:
            if source.kind != "NFC":
                continue
            outcome, other_fcs = audit.path_outcome(source)
            if audit.source_contains_fn(source):
                query_nfc_paths += 1
                if outcome.startswith("intended FC reached"):
                    query_nfc_intended += 1
                elif outcome == "other FC(s) only":
                    query_nfc_other_only += 1
                    for other in other_fcs:
                        other_transitions[f"{source.label}->{other}"] += 1
            elif outcome == "other FC(s) only":
                incidental_nfc_other_only += 1

    target = sum(audit.marker_overlaps_fn for audit in audits)
    linked_exact = sum(bool(audit.marker_linked_exact_records) for audit in audits)
    production_exact = sum(audit.marker_linked_production_exact for audit in audits)
    no_marker = sum(not audit.marker_records for audit in audits)
    ambiguous = sum(len(audit.sources) > 1 for audit in audits)
    query_linked = sum(bool(audit.query_linked_sources) for audit in audits)
    no_query_link = len(audits) - query_linked
    away = sum(bool(audit.marker_records) and not audit.marker_overlaps_fn for audit in audits)
    non_exact_at_target = target - linked_exact
    if production_exact:
        loss = "linked exact ALT survives production thresholds; review Cat4 assignment"
    elif no_query_link and target:
        loss = (
            f"mixed: {no_query_link}/{len(audits)} qnames have no extraction path tied "
            f"to an FN-containing FC and remain away from the FN; {target}/{len(audits)} "
            f"reach the target, including {linked_exact} exact ALT record(s) removed by "
            "MAPQ/base-quality thresholds"
        )
    elif no_query_link == len(audits):
        loss = (
            f"homologous-region assignment: all {len(audits)} qnames enter through "
            "FC paths not tied to the FN-containing FC and realign away from the FN"
        )
    elif target:
        details: list[str] = []
        if non_exact_at_target:
            details.append(
                f"{non_exact_at_target}/{len(audits)} target-reaching qnames do not encode the exact ALT"
            )
        if linked_exact:
            details.append(
                f"{linked_exact}/{len(audits)} encode the exact ALT but fail MAPQ/base-quality thresholds"
            )
        if no_marker:
            details.append(
                f"{no_marker}/{len(audits)} lose the ALT-haplotype-bearing record after realignment"
            )
        loss = "realignment representation: " + "; ".join(details)
    elif away:
        loss = (
            f"realignment placement: {away}/{len(audits)} ALT-haplotype records "
            "remain away from the FN coordinate"
        )
    elif no_marker == len(audits):
        loss = "realigned qname pairs no longer retain the verified ALT-haplotype sequence"
    else:
        loss = "ALT-haplotype reads realign outside the FN coordinate and outside catalogued FC intervals"

    return {
        "Assembly": assembly,
        "FN allele": site.label,
        "Eligible qnames": len(audits),
        "Qnames with query-linked extraction path": query_linked,
        "Qnames without query-linked extraction path": no_query_link,
        "Qnames eligible through NFC": len(nfc_audits),
        "Eligible NFC paths": len(nfc_paths),
        "Query-linked NFC paths": query_nfc_paths,
        "Query-linked NFC paths reaching intended FC": query_nfc_intended,
        "Query-linked NFC paths reaching only another FC": query_nfc_other_only,
        "Query-linked NFC-to-other-FC transitions": ";".join(
            f"{transition}(n={count})"
            for transition, count in sorted(other_transitions.items())
        ),
        "Incidental NFC paths reaching only another FC": incidental_nfc_other_only,
        "Qnames with ambiguous extraction provenance": ambiguous,
        "Qnames with ALT-haplotype record at FN anchor": target,
        "Qnames with ALT-haplotype record away from FN anchor": away,
        "Qnames with linked exact ALT before quality filters": linked_exact,
        "Qnames with linked exact ALT after MAPQ10/BQ15": production_exact,
        "Qnames with no ALT-haplotype record after realignment": no_marker,
        "Cat4 loss point": loss,
    }


def write_tsv(path: Path, rows: list[dict[str, object]]) -> None:
    if not rows:
        raise ValueError(f"refusing to write empty table: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), delimiter="\t")
        writer.writeheader()
        writer.writerows(rows)


def workbook_summary_text(original: object, summary: dict[str, object]) -> str:
    original_lines = [line for line in str(original or "").splitlines() if line]
    depth_lines = [line for line in original_lines if not line.startswith("Cause:") and not line.startswith("Cat4 path audit:")]
    transitions = summary["Query-linked NFC-to-other-FC transitions"] or "none"
    cause = (
        "Cause: Cat4 | " + str(summary["Cat4 loss point"])
        + "\nCat4 path audit: "
        + f"eligible qnames={summary['Eligible qnames']}; "
        + f"query-linked/unlinked={summary['Qnames with query-linked extraction path']}/"
        + f"{summary['Qnames without query-linked extraction path']}; "
        + f"NFC-eligible qnames={summary['Qnames eligible through NFC']}; "
        + f"query-linked NFC paths intended/other-only="
        + f"{summary['Query-linked NFC paths reaching intended FC']}/"
        + f"{summary['Query-linked NFC paths reaching only another FC']}; "
        + f"target-overlap/raw-exact/quality-pass={summary['Qnames with ALT-haplotype record at FN anchor']}/"
        + f"{summary['Qnames with linked exact ALT before quality filters']}/"
        + f"{summary['Qnames with linked exact ALT after MAPQ10/BQ15']}; "
        + f"NFC cross-FC={transitions}"
    )
    return "\n".join(depth_lines + [cause])


def save_workbook_atomic(workbook: openpyxl.Workbook, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(f".{output.name}.{os.getpid()}.tmp")
    workbook.save(temporary)
    os.replace(temporary, output)


def main() -> None:
    args = parse_args()
    workbook = openpyxl.load_workbook(args.workbook)
    if workbook.sheetnames != ["FN observations", "Qname mappings"]:
        raise ValueError(f"unexpected workbook sheets: {workbook.sheetnames}")
    observations = workbook["FN observations"]
    mappings = workbook["Qname mappings"]
    if worksheet_headers(observations) != OBSERVATION_COLUMNS:
        raise ValueError("FN observations schema differs from the agreed format")
    if worksheet_headers(mappings) != QNAME_COLUMNS:
        raise ValueError("Qname mappings schema differs from the agreed format")

    cat4_observation_rows: dict[tuple[str, str], int] = {}
    sites: dict[tuple[str, str], Site] = {}
    relevant_rgs: dict[tuple[str, str], set[str]] = defaultdict(set)
    for row_number, row in enumerate(
        observations.iter_rows(min_row=2, values_only=True), start=2
    ):
        if "Cause: Cat4 |" not in str(row[10]):
            continue
        assembly = str(row[0])
        site = parse_site(assembly, row[1])
        key = (assembly, site.label)
        cat4_observation_rows[key] = row_number
        sites[key] = site
        relevant_rgs[key].update(
            value for value in str(row[4]).splitlines() if RG_RE.fullmatch(value)
        )

    cat4_mapping_rows: dict[tuple[str, str, str], tuple[int, tuple[object, ...]]] = {}
    eligible_by_site: dict[tuple[str, str], set[str]] = defaultdict(set)
    sources_by_qname: dict[tuple[str, str, str], list[SourcePath]] = {}
    for row_number, row in enumerate(mappings.iter_rows(min_row=2, values_only=True), start=2):
        key = (str(row[0]), str(row[1]))
        if key not in cat4_observation_rows:
            continue
        qname = canonical_qname(str(row[2]))
        qkey = (*key, qname)
        cat4_mapping_rows[qkey] = (row_number, row)
        sources = parse_source_paths(row[7])
        if not sources:
            continue
        sources_by_qname[qkey] = sources
        eligible_by_site[key].add(qname)
        relevant_rgs[key].update(source.rg for source in sources)

    trace_by_qname: dict[tuple[str, str, str], list[TraceRecord]] = defaultdict(list)
    for assembly, haplotype_dir in HAPLOTYPE_DIRS.items():
        assembly_sites = {
            site_label for observed_assembly, site_label in sites if observed_assembly == assembly
        }
        rows = read_trace_rows(
            assembly,
            haplotype_dir / "pre_fp_input_qname_trace.tsv",
            assembly_sites,
        )
        trace_by_qname.update(rows)

    fc_catalogs: dict[str, dict[str, list[Interval]]] = {}
    for assembly in {key[0] for key in sites}:
        assembly_rgs = set().union(
            *(rgs for (observed_assembly, _), rgs in relevant_rgs.items() if observed_assembly == assembly)
        )
        fc_catalogs[assembly] = read_fc_catalog(PIPELINE_ROOTS[assembly], assembly_rgs)

    exact_by_site_qname: dict[tuple[str, str, str], list[ExactRecord]] = {}
    production_by_site_qname: dict[tuple[str, str, str], set[str]] = {}
    for key, site in sites.items():
        assembly, site_label = key
        exact, production = target_exact_records(
            assembly,
            site,
            relevant_rgs[key],
            eligible_by_site[key],
            args.samtools,
        )
        for qname in eligible_by_site[key]:
            exact_by_site_qname[(assembly, site_label, qname)] = exact.get(qname, [])
            production_by_site_qname[(assembly, site_label, qname)] = production.get(qname, set())

    audits_by_site: dict[tuple[str, str], list[QnameAudit]] = defaultdict(list)
    audits_by_qname: dict[tuple[str, str, str], QnameAudit] = {}
    for qkey, sources in sources_by_qname.items():
        assembly, site_label, qname = qkey
        _, mapping_row = cat4_mapping_rows[qkey]
        audit = QnameAudit(
            assembly=assembly,
            site=sites[(assembly, site_label)],
            qname=qname,
            evidence_source=str(mapping_row[3]),
            sources=sources,
            fastq_status=str(mapping_row[8]),
            trace_records=trace_by_qname.get(qkey, []),
            exact_records=exact_by_site_qname[qkey],
            production_exact_rgs=production_by_site_qname[qkey],
            fc_catalog=fc_catalogs[assembly],
        )
        audits_by_site[(assembly, site_label)].append(audit)
        audits_by_qname[qkey] = audit

    qname_rows = [
        qname_row(audit)
        for key in sorted(audits_by_site)
        for audit in sorted(audits_by_site[key], key=lambda value: value.qname)
    ]
    path_rows = [
        path_row(audit, source)
        for key in sorted(audits_by_site)
        for audit in sorted(audits_by_site[key], key=lambda value: value.qname)
        for source in audit.sources
    ]
    summary_rows = [
        allele_summary(assembly, sites[(assembly, site_label)], audits_by_site[(assembly, site_label)])
        for assembly, site_label in sorted(audits_by_site)
    ]

    qname_output = args.output_dir / "cat4_qname_summary.tsv"
    path_output = args.output_dir / "cat4_qname_extraction_paths.tsv"
    summary_output = args.output_dir / "cat4_allele_summary.tsv"
    write_tsv(qname_output, qname_rows)
    write_tsv(path_output, path_rows)
    write_tsv(summary_output, summary_rows)

    for summary in summary_rows:
        print(
            "ALLELE\t"
            + "\t".join(
                str(summary[column])
                for column in (
                    "Assembly",
                    "FN allele",
                    "Eligible qnames",
                    "Qnames with query-linked extraction path",
                    "Qnames without query-linked extraction path",
                    "Qnames eligible through NFC",
                    "Query-linked NFC paths reaching intended FC",
                    "Query-linked NFC paths reaching only another FC",
                    "Qnames with ALT-haplotype record at FN anchor",
                    "Qnames with linked exact ALT before quality filters",
                    "Qnames with linked exact ALT after MAPQ10/BQ15",
                    "Cat4 loss point",
                )
            )
        )

    expected_sites = 14
    expected_eligible_qnames = 220
    if len(summary_rows) != expected_sites:
        raise ValueError(f"expected {expected_sites} Cat4 sites, observed {len(summary_rows)}")
    if len(qname_rows) != expected_eligible_qnames:
        raise ValueError(
            f"expected {expected_eligible_qnames} eligible Cat4 qnames, observed {len(qname_rows)}"
        )
    if any("Yes |" not in audit.fastq_status for audit in audits_by_qname.values()):
        raise ValueError("an eligible Cat4 qname is absent from the extracted FASTQ")
    if any(not audit.trace_records for audit in audits_by_qname.values()):
        missing = [key for key, audit in audits_by_qname.items() if not audit.trace_records]
        raise ValueError(f"eligible Cat4 qnames missing from realigned trace: {missing[:5]}")

    print(f"cat4_sites={len(summary_rows)}")
    print(f"eligible_cat4_qnames={len(qname_rows)}")
    print(f"eligible_extraction_paths={len(path_rows)}")
    print(f"qname_output={qname_output.resolve()}")
    print(f"path_output={path_output.resolve()}")
    print(f"summary_output={summary_output.resolve()}")

    if not args.update_workbook:
        print("workbook_updated=false")
        return

    summary_by_site = {
        (str(summary["Assembly"]), str(summary["FN allele"])): summary
        for summary in summary_rows
    }
    for key, row_number in cat4_observation_rows.items():
        observations.cell(
            row_number,
            11,
            workbook_summary_text(observations.cell(row_number, 11).value, summary_by_site[key]),
        )
    for qkey, (row_number, row) in cat4_mapping_rows.items():
        audit = audits_by_qname.get(qkey)
        if audit is None:
            mappings.cell(row_number, 10, "Not extracted | not eligible")
        else:
            mappings.cell(row_number, 10, audit.assertion())

    output_workbook = args.output_workbook or args.workbook
    save_workbook_atomic(workbook, output_workbook)
    print(f"workbook_output={output_workbook.resolve()}")
    print("workbook_updated=true")


if __name__ == "__main__":
    main()
